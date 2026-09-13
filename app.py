import io, re, zipfile
from typing import List, Dict, Tuple

import numpy as np
import pandas as pd
import streamlit as st
from scipy import stats

st.set_page_config(page_title="Interference Analysis App", layout="wide")
st.title("Interference Analysis: Control vs Interference")
st.caption("Compare control vs interference conditions with mean/median shift, assumption checks, optional outlier removal, bootstrap CIs, and downloadable outputs.")

ID_HINTS = ["batch_id", "bloodSampleId", "bloodSampleID", "sampleId", "deviceId", "serialNumber", "patientUserId"]
DEFAULT_ANALYTE_HINTS = ["RBC", "WBC_2", "PLT_3", "HCT", "HGB", "MCV_3", "RDW_3", "MCH", "MCHC", "NEUT_2", "LYMPH_2", "MXD_2", "PLT", "MCV", "RDW"]
AUTO_OUTLIER_METHOD = "Automatic: Shapiro-Wilk -> Gcrit if normal, Robust MAD if non-normal"
GCRIT_OUTLIER_METHOD = "Gcrit Grubbs-like: remove largest |value-mean|/SD if >= Gcrit"
MAD_OUTLIER_METHOD = "Robust MAD modified-z: remove largest robust z if >= threshold"

# ------------------------- input helpers -------------------------
def read_upload(uploaded):
    name = uploaded.name.lower()
    if name.endswith(".csv"):
        return pd.read_csv(uploaded)
    return pd.read_excel(uploaded, engine="openpyxl")

def numeric_cols(df: pd.DataFrame) -> List[str]:
    cols = []
    for c in df.columns:
        if c in ID_HINTS:
            continue
        s = pd.to_numeric(df[c], errors="coerce")
        if s.notna().sum() >= 3:
            cols.append(c)
    return cols

def find_col(df, candidates):
    lower = {c.lower(): c for c in df.columns}
    for cand in candidates:
        if cand.lower() in lower:
            return lower[cand.lower()]
    return None

def extract_pair_key_from_sample(x: str, condition_value: str = "") -> str:
    s = str(x)
    m = re.search(r"[-_](?:C|I|CTRL|CONTROL|INT|INTERFERENCE)?\s*(\d+)\s*$", s, flags=re.I)
    if m:
        return m.group(1)
    nums = re.findall(r"(\d+)", s)
    return nums[-1] if nums else s

# ------------------------- stats helpers -------------------------
def mad_sd(x):
    x = np.asarray(pd.to_numeric(pd.Series(x), errors="coerce").dropna(), dtype=float)
    if len(x) == 0:
        return np.nan
    med = np.median(x)
    mad = np.median(np.abs(x - med))
    return 1.4826 * mad

def grubbs_gcrit(n: int, alpha: float = 0.01, tail: str = "Two-sided") -> float:
    """Return the classical Grubbs critical value for n observations.

    Two-sided uses t_(1-alpha/(2n), n-2); one-sided uses
    t_(1-alpha/n, n-2). This function is used only when the user
    selects automatic Gcrit calculation.
    """
    try:
        n = int(n)
        alpha = float(alpha)
    except Exception:
        return np.nan
    if n < 3 or not np.isfinite(alpha) or alpha <= 0 or alpha >= 1:
        return np.nan
    denom = 2 * n if str(tail).lower().startswith("two") else n
    tcrit = stats.t.ppf(1 - alpha / denom, df=n - 2)
    if not np.isfinite(tcrit):
        return np.nan
    return float(((n - 1) / np.sqrt(n)) * np.sqrt((tcrit ** 2) / (n - 2 + tcrit ** 2)))

def ci_quantiles(vals, alpha=0.05):
    vals = np.asarray(vals, dtype=float)
    vals = vals[np.isfinite(vals)]
    if len(vals) == 0:
        return np.nan, np.nan
    return np.percentile(vals, [100*alpha/2, 100*(1-alpha/2)])

def bootstrap_unpaired(x, y, stat_func, n_boot=2000, seed=1):
    rng = np.random.default_rng(seed)
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) < 2 or len(y) < 2:
        return np.nan, np.nan
    out = []
    for _ in range(int(n_boot)):
        xb = rng.choice(x, size=len(x), replace=True)
        yb = rng.choice(y, size=len(y), replace=True)
        out.append(stat_func(xb, yb))
    return ci_quantiles(out)

def bootstrap_paired(diffs, stat_func=np.mean, n_boot=2000, seed=1):
    rng = np.random.default_rng(seed)
    diffs = np.asarray(diffs, dtype=float)
    diffs = diffs[np.isfinite(diffs)]
    if len(diffs) < 2:
        return np.nan, np.nan
    out = []
    for _ in range(int(n_boot)):
        db = rng.choice(diffs, size=len(diffs), replace=True)
        out.append(stat_func(db))
    return ci_quantiles(out)

def permutation_pvalue_unpaired(x, y, n_perm=5000, seed=1):
    rng = np.random.default_rng(seed)
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) < 2 or len(y) < 2:
        return np.nan
    obs = abs(np.mean(y) - np.mean(x))
    z = np.concatenate([x, y])
    nx = len(x)
    count = 0
    for _ in range(int(n_perm)):
        zp = rng.permutation(z)
        stat = abs(np.mean(zp[nx:]) - np.mean(zp[:nx]))
        if stat >= obs:
            count += 1
    return (count + 1) / (n_perm + 1)

def benjamini_hochberg(pvals):
    p = np.asarray(pvals, dtype=float)
    q = np.full_like(p, np.nan, dtype=float)
    mask = np.isfinite(p)
    pm = p[mask]
    if len(pm) == 0:
        return q
    order = np.argsort(pm)
    ranked = pm[order]
    m = len(pm)
    adj = ranked * m / (np.arange(1, m+1))
    adj = np.minimum.accumulate(adj[::-1])[::-1]
    adj = np.minimum(adj, 1.0)
    qmask = np.empty_like(pm)
    qmask[order] = adj
    q[mask] = qmask
    return q

def assumption_checks(x, y):
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    residuals = np.concatenate([x - np.mean(x), y - np.mean(y)]) if len(x) and len(y) else np.array([])
    shapiro_p = stats.shapiro(residuals).pvalue if 3 <= len(residuals) <= 5000 else np.nan
    lev_mean_p = stats.levene(x, y, center="mean").pvalue if len(x) >= 2 and len(y) >= 2 else np.nan
    brown_p = stats.levene(x, y, center="median").pvalue if len(x) >= 2 and len(y) >= 2 else np.nan
    return shapiro_p, lev_mean_p, brown_p


def welch_mean_difference_ci(x, y, alpha=0.05):
    """Welch-Satterthwaite CI for mean(I)-mean(C)."""
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) < 2 or len(y) < 2:
        return np.nan, np.nan, np.nan, np.nan
    diff = float(np.mean(y) - np.mean(x))
    vx = float(np.var(x, ddof=1)); vy = float(np.var(y, ddof=1))
    se2 = vx/len(x) + vy/len(y)
    if not np.isfinite(se2) or se2 <= 0:
        return diff, np.nan, np.nan, np.nan
    se = np.sqrt(se2)
    df = se2**2 / (((vx/len(x))**2)/(len(x)-1) + ((vy/len(y))**2)/(len(y)-1))
    tcrit = stats.t.ppf(1-alpha/2, df)
    return diff, float(diff-tcrit*se), float(diff+tcrit*se), float(df)


def percent_shift_delta_ci(x, y, alpha=0.05):
    """Delta-method normal CI for 100*(mean(I)-mean(C))/mean(C)."""
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) < 2 or len(y) < 2:
        return np.nan, np.nan, np.nan
    mx, my = float(np.mean(x)), float(np.mean(y))
    if mx == 0:
        return np.nan, np.nan, np.nan
    est = 100.0*(my-mx)/mx
    vx = float(np.var(x, ddof=1))/len(x); vy = float(np.var(y, ddof=1))/len(y)
    var_g = (100.0*my/(mx**2))**2*vx + (100.0/mx)**2*vy
    if not np.isfinite(var_g) or var_g < 0:
        return est, np.nan, np.nan
    z = stats.norm.ppf(1-alpha/2)
    se = np.sqrt(var_g)
    return float(est), float(est-z*se), float(est+z*se)


def hodges_lehmann_shift(x, y, max_pairs=2000000, seed=1):
    """Independent-samples Hodges-Lehmann location shift: median of I-C pairwise differences."""
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) == 0 or len(y) == 0:
        return np.nan
    total = len(x)*len(y)
    if total <= int(max_pairs):
        return float(np.median((y[:,None]-x[None,:]).ravel()))
    rng = np.random.default_rng(seed)
    ix = rng.integers(0, len(x), size=int(max_pairs)); iy = rng.integers(0, len(y), size=int(max_pairs))
    return float(np.median(y[iy]-x[ix]))


def bootstrap_hodges_lehmann_ci(x, y, n_boot=2000, seed=1, alpha=0.05):
    rng = np.random.default_rng(seed)
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    if len(x) < 2 or len(y) < 2:
        return np.nan, np.nan
    vals=[]
    for _ in range(int(n_boot)):
        xb=rng.choice(x,len(x),replace=True); yb=rng.choice(y,len(y),replace=True)
        vals.append(hodges_lehmann_shift(xb,yb,max_pairs=200000,seed=int(rng.integers(1,2**31-1))))
    return ci_quantiles(vals, alpha=alpha)


def normalize_bool(value) -> bool:
    if pd.isna(value):
        return False
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, float, np.integer, np.floating)) and np.isfinite(value):
        return bool(int(value))
    return str(value).strip().lower() in {"true", "t", "1", "yes", "y", "flagged"}


def choose_automatic_outlier_method(x, y) -> str:
    """Use the existing normality diagnostic only to select the existing outlier rule."""
    shapiro_p, _, _ = assumption_checks(x, y)
    return GCRIT_OUTLIER_METHOD if np.isfinite(shapiro_p) and shapiro_p >= 0.05 else MAD_OUTLIER_METHOD


def choose_primary_outcome(row: Dict, alpha: float) -> Dict:
    """Select exactly one reportable branch from the normality diagnostic."""
    normal = row.get("residuals_normal", None) is True
    if normal:
        row["statistical_branch"] = "parametric"
        row["primary_test"] = "Welch t-test"
        row["selected_raw_p_value"] = row.get("welch_t_p_primary", np.nan)
        row["recommended_p_value"] = row["selected_raw_p_value"]
        row["selected_effect_estimate"] = row.get("mean_difference_I_minus_C", np.nan)
        row["selected_95CI_low"] = row.get("mean_difference_welch_95CI_low", np.nan)
        row["selected_95CI_high"] = row.get("mean_difference_welch_95CI_high", np.nan)
    else:
        row["statistical_branch"] = "nonparametric"
        row["primary_test"] = "Mann-Whitney U"
        row["selected_raw_p_value"] = row.get("mann_whitney_p_robust", np.nan)
        row["recommended_p_value"] = row["selected_raw_p_value"]
        row["selected_effect_estimate"] = row.get("hodges_lehmann_shift_I_minus_C", np.nan)
        row["selected_95CI_low"] = row.get("hodges_lehmann_shift_95CI_low", np.nan)
        row["selected_95CI_high"] = row.get("hodges_lehmann_shift_95CI_high", np.nan)
    p = row.get("selected_raw_p_value", np.nan)
    row["significant_at_alpha"] = bool(p < alpha) if np.isfinite(p) else None
    return row


def global_flag_audit(excluded: pd.DataFrame, analytes: List[str], cond_col: str, batch_col: str, sample_col: str, global_flag_col: str) -> pd.DataFrame:
    cols = ["batch_id", "condition", "analyte", "sample_id", "global_flag"]
    if excluded is None or excluded.empty:
        return pd.DataFrame(columns=cols)
    rows = []
    for _, r in excluded.iterrows():
        for analyte in analytes:
            if analyte in excluded.columns and pd.notna(r.get(analyte, np.nan)):
                rows.append({
                    "batch_id": _id_value(r, batch_col),
                    "condition": r.get(cond_col, ""),
                    "analyte": analyte,
                    "sample_id": _id_value(r, sample_col),
                    "global_flag": r.get(global_flag_col, True) if global_flag_col in excluded.columns else True,
                })
    return pd.DataFrame(rows, columns=cols)

# ------------------------- outlier helpers -------------------------
def _id_value(row, col):
    return row[col] if col and col != "None" and col in row.index else ""

def detect_outliers_one_condition(work: pd.DataFrame, analyte: str, condition_value: str,
                                  method: str, max_remove: int, gcrit: float,
                                  modified_z_threshold: float, robust_interval_z: float,
                                  batch_col: str, sample_col: str, device_col: str,
                                  scope_name: str, gcrit_mode: str = "Manual",
                                  gcrit_alpha: float = 0.01, gcrit_tail: str = "Two-sided") -> Tuple[List[int], List[Dict]]:
    """Detect up to max_remove outliers within one condition/analyte/scope. Returns original DataFrame indices."""
    sub = work.loc[work["__condition_str__"] == str(condition_value), [analyte]].copy()
    sub[analyte] = pd.to_numeric(sub[analyte], errors="coerce")
    sub = sub.dropna()
    remaining = list(sub.index)
    removed = []
    logs = []
    if method == "None" or int(max_remove) <= 0 or len(remaining) < 3:
        return removed, logs

    for step in range(int(max_remove)):
        vals = work.loc[remaining, analyte].astype(float)
        if vals.notna().sum() < 3:
            break
        x = vals.to_numpy(dtype=float)

        chosen_idx = None
        metric = np.nan
        threshold = np.nan
        direction = ""
        details = ""

        if method.startswith("Gcrit"):
            mu = float(np.mean(x)); sd = float(np.std(x, ddof=1))
            if not np.isfinite(sd) or sd == 0:
                break
            gvals = np.abs(x - mu) / sd
            k = int(np.argmax(gvals))
            metric = float(gvals[k])
            if str(gcrit_mode).lower().startswith("automatic"):
                threshold = grubbs_gcrit(len(x), alpha=float(gcrit_alpha), tail=str(gcrit_tail))
                threshold_label = f"automatic Gcrit={threshold:.4g}; n={len(x)}; alpha={float(gcrit_alpha):.4g}; tail={gcrit_tail}"
            else:
                threshold = float(gcrit)
                threshold_label = f"manual Gcrit={threshold:.4g}"
            if np.isfinite(threshold) and metric >= threshold:
                chosen_idx = remaining[k]
                direction = "high" if x[k] > mu else "low"
                details = f"G={metric:.4g}; mean={mu:.4g}; sd={sd:.4g}; {threshold_label}"
            else:
                break

        elif method.startswith("Robust MAD"):
            med = float(np.median(x)); mad = float(np.median(np.abs(x - med)))
            if not np.isfinite(mad) or mad == 0:
                break
            modz = 0.6745 * (x - med) / mad
            k = int(np.argmax(np.abs(modz)))
            metric = float(abs(modz[k])); threshold = float(modified_z_threshold)
            if metric >= modified_z_threshold:
                chosen_idx = remaining[k]
                direction = "high" if x[k] > med else "low"
                details = f"modified_z={modz[k]:.4g}; median={med:.4g}; MAD={mad:.4g}; threshold={modified_z_threshold}"
            else:
                break

        elif method.startswith("95% robust interval"):
            med = float(np.median(x)); rsd = mad_sd(x)
            if not np.isfinite(rsd) or rsd == 0:
                break
            lo = med - robust_interval_z * rsd
            hi = med + robust_interval_z * rsd
            distances = np.maximum(lo - x, x - hi)
            k = int(np.argmax(distances))
            metric = float(distances[k]); threshold = 0.0
            if metric > 0:
                chosen_idx = remaining[k]
                direction = "high" if x[k] > hi else "low"
                details = f"value outside robust interval [{lo:.4g}, {hi:.4g}]; median={med:.4g}; robust_SD={rsd:.4g}; z={robust_interval_z}"
            else:
                break

        if chosen_idx is None:
            break

        row = work.loc[chosen_idx]
        logs.append({
            "scope": scope_name,
            "condition": condition_value,
            "analyte": analyte,
            "removed_order": step + 1,
            "outlier_method": method,
            "row_index": int(chosen_idx) if isinstance(chosen_idx, (int, np.integer)) else str(chosen_idx),
            "batch_id": _id_value(row, batch_col),
            "sample_id": _id_value(row, sample_col),
            "device_id": _id_value(row, device_col),
            "value_removed": row[analyte],
            "direction": direction,
            "outlier_metric": metric,
            "outlier_threshold": threshold,
            "gcrit_mode": gcrit_mode if method.startswith("Gcrit") else "",
            "gcrit_alpha": gcrit_alpha if method.startswith("Gcrit") else "",
            "gcrit_tail": gcrit_tail if method.startswith("Gcrit") else "",
            "details": details,
        })
        removed.append(chosen_idx)
        remaining.remove(chosen_idx)

    return removed, logs

def apply_outlier_removal_for_analyte(work: pd.DataFrame, analyte: str, control_val: str, int_val: str,
                                      method: str, max_remove_per_condition: int, gcrit: float,
                                      modified_z_threshold: float, robust_interval_z: float,
                                      batch_col: str, sample_col: str, device_col: str,
                                      scope_name: str, gcrit_mode: str = "Manual",
                                      gcrit_alpha: float = 0.01, gcrit_tail: str = "Two-sided") -> Tuple[pd.DataFrame, pd.DataFrame]:
    work = work.copy()
    work["__condition_str__"] = work["__condition_str__"].astype(str)
    all_removed = []
    all_logs = []
    for cond in [str(control_val), str(int_val)]:
        removed, logs = detect_outliers_one_condition(
            work, analyte, cond, method, max_remove_per_condition, gcrit,
            modified_z_threshold, robust_interval_z, batch_col, sample_col, device_col, scope_name,
            gcrit_mode, gcrit_alpha, gcrit_tail
        )
        all_removed.extend(removed)
        all_logs.extend(logs)
    cleaned = work.drop(index=list(dict.fromkeys(all_removed)), errors="ignore").copy()
    return cleaned, pd.DataFrame(all_logs)

# ------------------------- comparison helpers -------------------------
def compare_unpaired(x, y, control_label, int_label, n_boot, do_boot, seed):
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    x = x[np.isfinite(x)]; y = y[np.isfinite(y)]
    row = {}
    row[f"n_{control_label}"] = len(x); row[f"n_{int_label}"] = len(y)
    row[f"mean_{control_label}"] = np.mean(x) if len(x) else np.nan
    row[f"sd_{control_label}"] = np.std(x, ddof=1) if len(x) > 1 else np.nan
    row[f"median_{control_label}"] = np.median(x) if len(x) else np.nan
    row[f"iqr_{control_label}"] = np.subtract(*np.percentile(x, [75,25])) if len(x) else np.nan
    row[f"mean_{int_label}"] = np.mean(y) if len(y) else np.nan
    row[f"sd_{int_label}"] = np.std(y, ddof=1) if len(y) > 1 else np.nan
    row[f"median_{int_label}"] = np.median(y) if len(y) else np.nan
    row[f"iqr_{int_label}"] = np.subtract(*np.percentile(y, [75,25])) if len(y) else np.nan
    row["mean_difference_I_minus_C"] = row[f"mean_{int_label}"] - row[f"mean_{control_label}"]
    row["median_difference_I_minus_C"] = row[f"median_{int_label}"] - row[f"median_{control_label}"]
    row["percent_shift_mean"] = 100 * row["mean_difference_I_minus_C"] / row[f"mean_{control_label}"] if row[f"mean_{control_label}"] != 0 else np.nan
    row["percent_shift_median"] = 100 * row["median_difference_I_minus_C"] / row[f"median_{control_label}"] if row[f"median_{control_label}"] != 0 else np.nan
    row["mad_sd_control"] = mad_sd(x)
    row["mad_sd_interference"] = mad_sd(y)
    sh, lev, bf = assumption_checks(x, y)
    row["shapiro_wilk_p_residuals"] = sh
    row["levene_mean_p"] = lev
    row["brown_forsythe_median_p"] = bf
    row["residuals_normal"] = bool(sh >= 0.05) if np.isfinite(sh) else None
    row["equal_variance_levene"] = bool(lev >= 0.05) if np.isfinite(lev) else None
    row["equal_variance_brown_forsythe"] = bool(bf >= 0.05) if np.isfinite(bf) else None
    try: row["student_t_p_equal_var"] = stats.ttest_ind(x, y, equal_var=True).pvalue
    except Exception: row["student_t_p_equal_var"] = np.nan
    try: row["welch_t_p_primary"] = stats.ttest_ind(x, y, equal_var=False).pvalue
    except Exception: row["welch_t_p_primary"] = np.nan
    try: row["mann_whitney_p_robust"] = stats.mannwhitneyu(x, y, alternative="two-sided").pvalue
    except Exception: row["mann_whitney_p_robust"] = np.nan
    row["permutation_p_mean_diff"] = permutation_pvalue_unpaired(x, y, seed=seed) if len(x) >= 2 and len(y) >= 2 else np.nan
    _, wlo, whi, wdf = welch_mean_difference_ci(x, y)
    row["mean_difference_welch_95CI_low"] = wlo
    row["mean_difference_welch_95CI_high"] = whi
    row["welch_satterthwaite_df"] = wdf
    _, plo, phi = percent_shift_delta_ci(x, y)
    row["percent_shift_mean_normal_95CI_low"] = plo
    row["percent_shift_mean_normal_95CI_high"] = phi
    row["hodges_lehmann_shift_I_minus_C"] = hodges_lehmann_shift(x, y, seed=seed)
    row["hodges_lehmann_shift_95CI_low"] = np.nan
    row["hodges_lehmann_shift_95CI_high"] = np.nan
    row["mean_difference_95CI_low"] = np.nan; row["mean_difference_95CI_high"] = np.nan
    row["percent_shift_mean_95CI_low"] = np.nan; row["percent_shift_mean_95CI_high"] = np.nan
    row["median_difference_95CI_low"] = np.nan; row["median_difference_95CI_high"] = np.nan
    row["percent_shift_median_95CI_low"] = np.nan; row["percent_shift_median_95CI_high"] = np.nan
    if do_boot:
        hlo, hhi = bootstrap_hodges_lehmann_ci(x, y, n_boot=n_boot, seed=seed+77)
        row["hodges_lehmann_shift_95CI_low"] = hlo
        row["hodges_lehmann_shift_95CI_high"] = hhi
        lo, hi = bootstrap_unpaired(x, y, lambda xb,yb: np.mean(yb)-np.mean(xb), n_boot=n_boot, seed=seed)
        row["mean_difference_95CI_low"] = lo; row["mean_difference_95CI_high"] = hi
        lo, hi = bootstrap_unpaired(x, y, lambda xb,yb: 100*(np.mean(yb)-np.mean(xb))/np.mean(xb) if np.mean(xb)!=0 else np.nan, n_boot=n_boot, seed=seed+11)
        row["percent_shift_mean_95CI_low"] = lo; row["percent_shift_mean_95CI_high"] = hi
        lo, hi = bootstrap_unpaired(x, y, lambda xb,yb: np.median(yb)-np.median(xb), n_boot=n_boot, seed=seed+22)
        row["median_difference_95CI_low"] = lo; row["median_difference_95CI_high"] = hi
        lo, hi = bootstrap_unpaired(x, y, lambda xb,yb: 100*(np.median(yb)-np.median(xb))/np.median(xb) if np.median(xb)!=0 else np.nan, n_boot=n_boot, seed=seed+33)
        row["percent_shift_median_95CI_low"] = lo; row["percent_shift_median_95CI_high"] = hi
    return row

def compare_paired(df_sub, analyte, cond_col, control_val, int_val, pair_cols, n_boot, do_boot, seed):
    work = df_sub[pair_cols + [cond_col, analyte]].copy()
    work[analyte] = pd.to_numeric(work[analyte], errors="coerce")
    wide = work.pivot_table(index=pair_cols, columns=cond_col, values=analyte, aggfunc="mean")
    if control_val not in wide.columns or int_val not in wide.columns:
        return None, pd.DataFrame()
    paired = wide[[control_val, int_val]].dropna().reset_index()
    x = paired[control_val].to_numpy(dtype=float)
    y = paired[int_val].to_numpy(dtype=float)
    diffs = y - x
    row = compare_unpaired(x, y, str(control_val), str(int_val), n_boot, do_boot, seed)
    row["paired_n"] = len(diffs)
    row["paired_mean_difference_I_minus_C"] = np.mean(diffs) if len(diffs) else np.nan
    row["paired_percent_shift_mean"] = 100*np.mean(diffs)/np.mean(x) if len(diffs) and np.mean(x)!=0 else np.nan
    try: row["paired_t_p"] = stats.ttest_rel(x, y).pvalue
    except Exception: row["paired_t_p"] = np.nan
    try: row["wilcoxon_p"] = stats.wilcoxon(diffs).pvalue if len(diffs) >= 2 else np.nan
    except Exception: row["wilcoxon_p"] = np.nan
    if do_boot and len(diffs) >= 2:
        lo, hi = bootstrap_paired(diffs, np.mean, n_boot, seed+44)
        row["paired_mean_difference_95CI_low"] = lo; row["paired_mean_difference_95CI_high"] = hi
    return row, paired

INTERFERENCE_SUMMARY_COLUMNS = [
    "scope", "result_type", "analyte", "n_C", "n_I", "Shapiro-Wilk\np value",
    "Are the data\nnormally distributed?", "Statistical Test Applied", "Recommended p value",
    "BH-FDR corrected p value (q)", "Statistically significant?", "mean_C", "sd_C", "median_C", "iqr_C",
    "mean_I", "sd_I", "median_I", "iqr_I", "mean_difference_I_minus_C", "median_difference_I_minus_C",
    "percent_shift_mean", "percent_shift_median", "mad_sd_control", "mad_sd_interference", "levene_mean_p",
    "brown_forsythe_median_p", "equal_variance_levene", "equal_variance_brown_forsythe", "student_t_p_equal_var",
    "welch_t_p_primary", "mann_whitney_p_robust", "permutation_p_mean_diff", "mean_difference_welch_95CI_low",
    "mean_difference_welch_95CI_high", "welch_satterthwaite_df", "percent_shift_mean_normal_95CI_low",
    "percent_shift_mean_normal_95CI_high", "hodges_lehmann_shift_I_minus_C", "control_condition",
    "interference_condition", "outlier_method", "n_outliers_removed",
    "selected_effect_estimate", "selected_95CI_low", "selected_95CI_high", "welch_t_q_BH_FDR_within_table",
    "mann_whitney_q_BH_FDR_within_table", "final_multiple_testing_decision",
]
OUTLIER_LOG_COLUMNS = [
    "scope", "condition", "analyte", "removed_order", "outlier_method", "row_index", "batch_id", "sample_id",
    "device_id", "value_removed", "direction", "outlier_metric", "outlier_threshold", "gcrit_mode", "gcrit_alpha",
    "gcrit_tail", "details",
]
GLOBAL_FLAG_COLUMNS = ["batch_id", "condition", "analyte", "sample_id", "global_flag"]


def finalize_selected_fdr(tbl: pd.DataFrame, alpha: float) -> pd.DataFrame:
    if tbl is None or tbl.empty:
        return tbl.copy() if isinstance(tbl, pd.DataFrame) else pd.DataFrame()
    out = tbl.copy()
    out["welch_t_q_BH_FDR_within_table"] = benjamini_hochberg(out["welch_t_p_primary"].values)
    out["mann_whitney_q_BH_FDR_within_table"] = benjamini_hochberg(out["mann_whitney_p_robust"].values)
    out["selected_q_BH_FDR"] = np.where(
        out["residuals_normal"].eq(True),
        out["welch_t_q_BH_FDR_within_table"],
        out["mann_whitney_q_BH_FDR_within_table"],
    )
    out["significant_after_BH_FDR"] = out["selected_q_BH_FDR"] < float(alpha)
    out["final_multiple_testing_decision"] = np.where(
        out["significant_after_BH_FDR"],
        "Significant interference detected",
        "No statistically significant interference detected",
    )
    return out


def format_interference_summary(tbl: pd.DataFrame, control_val: str, int_val: str, alpha: float) -> pd.DataFrame:
    """Match the supplied interference_results_example workbook column-for-column."""
    if tbl is None or tbl.empty:
        return pd.DataFrame(columns=INTERFERENCE_SUMMARY_COLUMNS)
    out = pd.DataFrame(index=tbl.index)
    get = lambda c: tbl[c] if c in tbl.columns else pd.Series(np.nan, index=tbl.index)
    out["scope"] = get("scope")
    out["result_type"] = get("result_type")
    out["analyte"] = get("analyte")
    out["n_C"] = get(f"n_{control_val}")
    out["n_I"] = get(f"n_{int_val}")
    out["Shapiro-Wilk\np value"] = get("shapiro_wilk_p_residuals")
    out["Are the data\nnormally distributed?"] = get("residuals_normal")
    out["Statistical Test Applied"] = get("primary_test")
    # Report the one p-value selected by the normality-guided inferential branch,
    # matching the TR report structure (Test -> p value -> significance).
    out["Recommended p value"] = get("recommended_p_value")
    # Also surface the Benjamini-Hochberg FDR-adjusted value for the same selected
    # branch. Method-specific p/q values remain later in the row for auditability.
    out["BH-FDR corrected p value (q)"] = get("selected_q_BH_FDR")
    # TR-style significance is based on the reported/recommended selected p-value.
    # The BH-FDR decision is retained later in final_multiple_testing_decision.
    out["Statistically significant?"] = get("recommended_p_value").apply(
        lambda x: "YES" if pd.notna(x) and float(x) < float(alpha) else ("No" if pd.notna(x) else None)
    )
    for prefix, source in [("mean_C",f"mean_{control_val}"),("sd_C",f"sd_{control_val}"),("median_C",f"median_{control_val}"),("iqr_C",f"iqr_{control_val}"),
                           ("mean_I",f"mean_{int_val}"),("sd_I",f"sd_{int_val}"),("median_I",f"median_{int_val}"),("iqr_I",f"iqr_{int_val}")]:
        out[prefix] = get(source)
    passthrough = [
        "mean_difference_I_minus_C", "median_difference_I_minus_C", "percent_shift_mean", "percent_shift_median",
        "mad_sd_control", "mad_sd_interference", "levene_mean_p", "brown_forsythe_median_p", "equal_variance_levene",
        "equal_variance_brown_forsythe", "student_t_p_equal_var", "welch_t_p_primary", "mann_whitney_p_robust",
        "permutation_p_mean_diff", "mean_difference_welch_95CI_low", "mean_difference_welch_95CI_high",
        "welch_satterthwaite_df", "percent_shift_mean_normal_95CI_low", "percent_shift_mean_normal_95CI_high",
        "hodges_lehmann_shift_I_minus_C", "control_condition", "interference_condition", "outlier_method",
        "n_outliers_removed", "selected_effect_estimate",
        "selected_95CI_low", "selected_95CI_high", "welch_t_q_BH_FDR_within_table",
        "mann_whitney_q_BH_FDR_within_table", "final_multiple_testing_decision",
    ]
    for c in passthrough:
        out[c] = get(c)
    return out[INTERFERENCE_SUMMARY_COLUMNS].reset_index(drop=True)


def make_excel_output(result_tables: Dict[str, pd.DataFrame]) -> bytes:
    """Return the exact five-sheet single workbook requested in the supplied example."""
    buf = io.BytesIO()
    sheet_order = [
        "interference_summary_raw", "interference_summary_cleaned", "outlier_log", "global flag TRUE", "condition_device_counts"
    ]
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name in sheet_order:
            table = result_tables.get(name, pd.DataFrame())
            table.to_excel(writer, sheet_name=name, index=False)

        from openpyxl.styles import Font, Alignment
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row and ws.max_column:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for cells in ws.columns:
                letter = cells[0].column_letter
                width = max((len(str(c.value)) if c.value is not None else 0 for c in cells[:min(ws.max_row, 200)]), default=8) + 2
                ws.column_dimensions[letter].width = min(max(width, 10), 42)
    return buf.getvalue()


def run_interference_analysis(df_eligible: pd.DataFrame, df_global_excluded: pd.DataFrame, *,
                              cond_col: str, control_val: str, int_val: str, device_col: str,
                              batch_col: str, sample_col: str, global_flag_col: str,
                              analytes: List[str], device_mode: str, paired_mode: bool,
                              do_boot: bool, n_boot: int, alpha: float, seed: int,
                              outlier_method: str, max_remove_per_condition: int,
                              gcrit_mode: str, gcrit: float, gcrit_alpha: float, gcrit_tail: str,
                              modified_z_threshold: float, robust_interval_z: float) -> Dict[str, object]:
    """Pure analysis runner used by the UI and regression tests."""
    df2 = df_eligible.copy()
    df2[cond_col] = df2[cond_col].astype(str)
    df2 = df2[df2[cond_col].isin([str(control_val), str(int_val)])].copy()
    df2["__condition_str__"] = df2[cond_col].astype(str)
    if device_col == "None":
        df2["__device__"] = "pooled"
        device_col_use = "__device__"
    else:
        device_col_use = device_col
        df2[device_col_use] = df2[device_col_use].astype(str)

    if sample_col != "None":
        df2["__pair_key__"] = df2.apply(lambda r: extract_pair_key_from_sample(r[sample_col], r[cond_col]), axis=1)
    else:
        df2["__pair_key__"] = df2.groupby(cond_col).cumcount().astype(str)

    scopes = [("pooled_all_devices", df2.copy())]
    if device_mode.startswith("Analyze each"):
        for dev, sdf in df2.groupby(device_col_use):
            scopes.append((f"device_{dev}", sdf.copy()))

    summary_raw_rows, summary_clean_rows, outlier_logs = [], [], []
    paired_raw_rows, paired_clean_rows = [], []
    for scope_name, sdf in scopes:
        for analyte in analytes:
            cols = [cond_col, "__condition_str__", analyte, device_col_use, "__pair_key__"]
            for extra in [batch_col, sample_col]:
                if extra != "None" and extra in sdf.columns and extra not in cols:
                    cols.append(extra)
            work = sdf[cols].copy()
            work[analyte] = pd.to_numeric(work[analyte], errors="coerce")

            x_raw = work.loc[work["__condition_str__"] == str(control_val), analyte].dropna().values
            y_raw = work.loc[work["__condition_str__"] == str(int_val), analyte].dropna().values
            row_raw = compare_unpaired(x_raw, y_raw, str(control_val), str(int_val), int(n_boot), bool(do_boot), int(seed))
            row_raw.update({"scope": scope_name, "analyte": analyte, "result_type": "raw_no_outlier_removal",
                            "control_condition": control_val, "interference_condition": int_val,
                            "outlier_method": "None", "n_outliers_removed": 0})
            summary_raw_rows.append(choose_primary_outcome(row_raw, float(alpha)))

            actual_outlier_method = outlier_method
            if outlier_method == AUTO_OUTLIER_METHOD:
                actual_outlier_method = choose_automatic_outlier_method(x_raw, y_raw)
            cleaned = work.copy(); log_df = pd.DataFrame()
            if actual_outlier_method != "None" and int(max_remove_per_condition) > 0:
                cleaned, log_df = apply_outlier_removal_for_analyte(
                    work, analyte, control_val, int_val, actual_outlier_method, int(max_remove_per_condition), float(gcrit),
                    float(modified_z_threshold), float(robust_interval_z), batch_col, sample_col, device_col_use, scope_name,
                    gcrit_mode, float(gcrit_alpha), gcrit_tail
                )
                if not log_df.empty:
                    outlier_logs.append(log_df)

            x_clean = cleaned.loc[cleaned["__condition_str__"] == str(control_val), analyte].dropna().values
            y_clean = cleaned.loc[cleaned["__condition_str__"] == str(int_val), analyte].dropna().values
            row_clean = compare_unpaired(x_clean, y_clean, str(control_val), str(int_val), int(n_boot), bool(do_boot), int(seed))
            row_clean.update({"scope": scope_name, "analyte": analyte, "result_type": "cleaned_after_outlier_rule",
                              "control_condition": control_val, "interference_condition": int_val,
                              "outlier_method": actual_outlier_method,
                              "n_outliers_removed": int(len(work) - len(cleaned))})
            summary_clean_rows.append(choose_primary_outcome(row_clean, float(alpha)))

            if paired_mode:
                pair_cols = ["__pair_key__"] if scope_name == "pooled_all_devices" else ["__pair_key__", device_col_use]
                prow, _ = compare_paired(work, analyte, cond_col, str(control_val), str(int_val), pair_cols, int(n_boot), bool(do_boot), int(seed))
                if prow is not None:
                    prow.update({"scope": scope_name, "analyte": analyte, "result_type": "raw_no_outlier_removal"})
                    paired_raw_rows.append(prow)
                pcrow, _ = compare_paired(cleaned, analyte, cond_col, str(control_val), str(int_val), pair_cols, int(n_boot), bool(do_boot), int(seed))
                if pcrow is not None:
                    pcrow.update({"scope": scope_name, "analyte": analyte, "result_type": "cleaned_after_outlier_rule"})
                    paired_clean_rows.append(pcrow)

    summary_raw_internal = finalize_selected_fdr(pd.DataFrame(summary_raw_rows), float(alpha))
    summary_clean_internal = finalize_selected_fdr(pd.DataFrame(summary_clean_rows), float(alpha))
    summary_raw = format_interference_summary(summary_raw_internal, str(control_val), str(int_val), float(alpha))
    summary_clean = format_interference_summary(summary_clean_internal, str(control_val), str(int_val), float(alpha))

    outlier_log = pd.concat(outlier_logs, ignore_index=True) if outlier_logs else pd.DataFrame(columns=OUTLIER_LOG_COLUMNS)
    for c in OUTLIER_LOG_COLUMNS:
        if c not in outlier_log.columns: outlier_log[c] = np.nan
    outlier_log = outlier_log[OUTLIER_LOG_COLUMNS]

    excluded_for_conditions = df_global_excluded[df_global_excluded[cond_col].astype(str).isin([str(control_val), str(int_val)])].copy() if not df_global_excluded.empty else df_global_excluded.copy()
    global_log = global_flag_audit(excluded_for_conditions, analytes, cond_col, batch_col, sample_col, global_flag_col if global_flag_col != "None" else "")
    for c in GLOBAL_FLAG_COLUMNS:
        if c not in global_log.columns: global_log[c] = np.nan
    global_log = global_log[GLOBAL_FLAG_COLUMNS]
    counts = df2.groupby([cond_col, device_col_use], dropna=False).size().reset_index(name="n_rows")

    return {
        "interference_summary_raw": summary_raw,
        "interference_summary_cleaned": summary_clean,
        "outlier_log": outlier_log,
        "global flag TRUE": global_log,
        "condition_device_counts": counts,
        "paired_summary_raw": pd.DataFrame(paired_raw_rows),
        "paired_summary_cleaned": pd.DataFrame(paired_clean_rows),
    }


# ------------------------- UI -------------------------
uploaded = st.file_uploader("Upload Excel or CSV", type=["xlsx", "xls", "csv"])
if uploaded is None:
    st.info("Upload your interference file to begin. Expected design: condition column with control/interference rows, optional deviceId, and selected numeric analytes.")
    st.stop()

try:
    df = read_upload(uploaded)
except Exception as e:
    st.error(f"Could not read file: {e}")
    st.stop()

st.subheader("1) Data preview and detected columns")
st.write(f"Rows: **{len(df)}** | Columns: **{len(df.columns)}**")
st.dataframe(df.head(30), use_container_width=True)

condition_guess = find_col(df, ["Condition", "condition", "Group", "group"])
device_guess = find_col(df, ["deviceId", "Device", "device", "serialNumber"])
batch_guess = find_col(df, ["batch_id", "batchId", "Batch"])
sample_guess = find_col(df, ["bloodSampleId", "sampleId", "SampleID", "sample"])

c1, c2, c3, c4 = st.columns(4)
with c1:
    cond_col = st.selectbox("Condition column", options=list(df.columns), index=list(df.columns).index(condition_guess) if condition_guess in df.columns else 0)
with c2:
    device_col = st.selectbox("Device column (optional)", options=["None"] + list(df.columns), index=(["None"] + list(df.columns)).index(device_guess) if device_guess in df.columns else 0)
with c3:
    batch_col = st.selectbox("Batch ID column", options=["None"] + list(df.columns), index=(["None"] + list(df.columns)).index(batch_guess) if batch_guess in df.columns else 0)
with c4:
    sample_col = st.selectbox("Sample/replicate ID column", options=["None"] + list(df.columns), index=(["None"] + list(df.columns)).index(sample_guess) if sample_guess in df.columns else 0)

flag_options = ["None"] + list(df.columns)
global_guess = next((c for c in df.columns if c.lower() == "global_flag"), None)
g1, g2 = st.columns([2, 3])
with g1:
    global_flag_col = st.selectbox(
        "Global flag column", flag_options,
        index=flag_options.index(global_guess) if global_guess in flag_options else 0,
    )
with g2:
    treat_all_global_false = st.checkbox(
        "Treat all rows as global_flag = FALSE when no flag column is selected", value=False
    )

if global_flag_col != "None" and not treat_all_global_false:
    global_mask = df[global_flag_col].map(normalize_bool).fillna(False).astype(bool)
else:
    global_mask = pd.Series(False, index=df.index)
df_eligible = df.loc[~global_mask].copy()
df_global_excluded = df.loc[global_mask].copy()

conditions = sorted([str(x) for x in df_eligible[cond_col].dropna().unique()])
if len(conditions) < 2:
    st.error("Need at least two condition values, e.g. C and I.")
    st.stop()

c1, c2 = st.columns(2)
with c1:
    control_val = st.selectbox("Control/reference condition", options=conditions, index=0)
with c2:
    int_val = st.selectbox("Interference/test condition", options=conditions, index=1 if len(conditions)>1 else 0)

numcols = numeric_cols(df_eligible)
# Put the same requested analytes first as in the Short-Term app, while retaining
# every other numeric column as an optional choice.
def_default = [c for c in DEFAULT_ANALYTE_HINTS if c in numcols]
ordered_numcols = def_default + [c for c in numcols if c not in def_default]
if not def_default:
    def_default = ordered_numcols[:10]
analytes = st.multiselect(
    "Analyte columns to compare",
    options=ordered_numcols,
    default=def_default,
    help="Select exactly which analytes to analyze. Suggested analytes are ordered the same way as in the Short-Term app.",
)

st.subheader("2) Analysis settings")
c1, c2, c3, c4 = st.columns(4)
with c1:
    device_mode = st.selectbox("Device handling", ["Pool all devices", "Analyze each device separately + pooled"])
with c2:
    paired_mode = st.checkbox("Try paired analysis using sample key and/or device", value=False)
with c3:
    do_boot = st.checkbox("Bootstrap 95% CIs", value=False)
with c4:
    n_boot = st.number_input("Bootstrap iterations", min_value=200, max_value=20000, value=2000, step=200)

c1, c2 = st.columns(2)
with c1:
    alpha = st.number_input("Significance alpha", min_value=0.001, max_value=0.2, value=0.05, step=0.01)
with c2:
    seed = st.number_input("Random seed", min_value=1, max_value=999999, value=123, step=1)

st.subheader("3) Optional outlier sensitivity analysis")
outlier_method = st.selectbox(
    "Outlier method for cleaned/sensitivity results",
    [
        AUTO_OUTLIER_METHOD,
        "None",
        GCRIT_OUTLIER_METHOD,
        MAD_OUTLIER_METHOD,
        "95% robust interval: remove most extreme outside median ± z*MAD_SD",
    ],
    index=0,
)
c1, c2, c3, c4 = st.columns(4)
with c1:
    max_remove_per_condition = st.selectbox("Max outliers to remove per condition/analyte/scope", [0, 1, 2], index=1)
with c2:
    gcrit_mode = st.selectbox("Gcrit mode", ["Manual", "Automatic from n, alpha, tail"], index=1)
with c3:
    gcrit = st.number_input("Manual Gcrit value", min_value=0.0, value=3.135, step=0.001, format="%.3f")
with c4:
    gcrit_alpha = st.number_input("Automatic Grubbs alpha", min_value=0.0001, max_value=0.2, value=0.01, step=0.001, format="%.4f")

c1, c2, c3 = st.columns(3)
with c1:
    gcrit_tail = st.selectbox("Automatic Grubbs tail", ["Two-sided", "One-sided"], index=0)
with c2:
    modified_z_threshold = st.number_input("MAD modified-z threshold", min_value=0.1, value=3.5, step=0.1)
with c3:
    robust_interval_z = st.number_input("Robust interval z", min_value=0.5, value=1.96, step=0.01)

if outlier_method.startswith("Gcrit") and gcrit_mode.startswith("Automatic"):
    st.caption("Automatic Gcrit uses the classical Grubbs critical value with n recalculated within each condition × analyte × scope after each sequential removal. Two-sided uses t_(1-alpha/(2n), n-2); one-sided uses t_(1-alpha/n, n-2).")

st.markdown("""
**Single inferential outcome:** Shapiro-Wilk residual normality selects the existing branch: normal residuals → Welch two-sample t-test; non-normal/not-testable residuals → Mann-Whitney U.  
**Assumption checks:** Shapiro-Wilk on residuals, classic Levene, and Brown-Forsythe/median-centered Levene.  
**Outlier outputs:** the app retains the raw result as an audit baseline and reports one cleaned result after the automatically selected outlier rule, plus a separate outlier log. It does not export competing sensitivity/statistical-choice tables.
""")

if not analytes:
    st.warning("Select at least one analyte.")
    st.stop()

run = st.button("Run interference analysis", type="primary")
if not run:
    st.stop()

# Run the fully automated branch-selection pipeline.
analysis = run_interference_analysis(
    df_eligible, df_global_excluded,
    cond_col=cond_col, control_val=str(control_val), int_val=str(int_val), device_col=device_col,
    batch_col=batch_col, sample_col=sample_col, global_flag_col=global_flag_col, analytes=analytes,
    device_mode=device_mode, paired_mode=paired_mode, do_boot=do_boot, n_boot=int(n_boot),
    alpha=float(alpha), seed=int(seed), outlier_method=outlier_method,
    max_remove_per_condition=int(max_remove_per_condition), gcrit_mode=gcrit_mode, gcrit=float(gcrit),
    gcrit_alpha=float(gcrit_alpha), gcrit_tail=gcrit_tail, modified_z_threshold=float(modified_z_threshold),
    robust_interval_z=float(robust_interval_z),
)
summary_raw = analysis["interference_summary_raw"]
summary_clean = analysis["interference_summary_cleaned"]
outlier_log = analysis["outlier_log"]
global_flag_log = analysis["global flag TRUE"]
counts = analysis["condition_device_counts"]
paired_raw = analysis["paired_summary_raw"]
paired_clean = analysis["paired_summary_cleaned"]

st.subheader("4) Results: raw pooled/all devices primary table")
st.dataframe(summary_raw[summary_raw["scope"] == "pooled_all_devices"], use_container_width=True)

st.subheader("5) Results: cleaned/outlier-sensitivity pooled/all devices primary table")
st.dataframe(summary_clean[summary_clean["scope"] == "pooled_all_devices"], use_container_width=True)

st.subheader("6) Outlier log")
if outlier_log.empty:
    st.info("No outliers removed, or outlier removal was set to None/0.")
else:
    st.dataframe(outlier_log, use_container_width=True)

st.subheader("7) Optional paired diagnostic, if enabled")
if paired_raw.empty and paired_clean.empty:
    st.info("No paired result could be created, or paired analysis was off.")
else:
    st.write("Raw paired results")
    st.dataframe(paired_raw, use_container_width=True)
    st.write("Cleaned paired results")
    st.dataframe(paired_clean, use_container_width=True)

st.subheader("8) Condition/device counts")
st.dataframe(counts, use_container_width=True)

st.subheader("9) global_flag=TRUE exclusions")
if global_flag_log.empty:
    st.info("No rows were excluded by global_flag, or all rows were explicitly treated as global_flag = FALSE.")
else:
    st.dataframe(global_flag_log, use_container_width=True)

result_tables = {
    "interference_summary_raw": summary_raw,
    "interference_summary_cleaned": summary_clean,
    "outlier_log": outlier_log,
    "global flag TRUE": global_flag_log,
    "condition_device_counts": counts,
}
excel_bytes = make_excel_output(result_tables)

st.download_button(
    "Download combined Excel results",
    excel_bytes,
    "interference_results.xlsx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
st.success("Done. Raw and cleaned summaries now show the selected test, Recommended p value, BH-FDR corrected p value (q), and TR-style statistical significance as the primary report columns.")
